"""Collaudo del confronto fra polizze (confronto_polizze.py).

Le polizze sono i due PDF di prova veri (polizze_prova.py), letti con il lettore PDF vero
di A'PLAUD'I: le tabelle arrivano con una cella per riga, come nei PDF reali.
Simulato SOLO il modello: le sue risposte contengono di proposito gli errori che un
modello fa (citazione inventata, numero sbagliato, pagina sbagliata, deroga dichiarata a
torto, frase con un numero inventato, giudizio incoerente). Il codice li deve fermare tutti.

Esecuzione:  python -m pytest -q test_confronto_polizze.py
"""
import io
import json

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

import confronto_polizze as cp
import documenti_aziendali as docs
import polizze_prova


# ------------------------------------------------------------------ le risposte del modello finto
def V(chiave, nome, valore, citazione, pagina, deroga=False, inclusa=None):
    return {"chiave": chiave, "nome": nome, "valore": valore, "citazione": citazione,
            "pagina": pagina, "deroga": deroga, "inclusa": inclusa}


LETTURA_A = {"ramo": "vita - temporanea caso morte", "compagnia": "VERSILIA VITA S.p.A.",
             "prodotto": "VERSILIA PROTEZIONE 200", "voci": [
    V("prodotto", "Tipo di contratto", "temporanea caso morte a capitale costante",
      "Contratto di assicurazione temporanea caso morte a capitale costante", "A1"),
    V("premio", "Premio annuo lordo", "348,00 EUR annui", "Premio annuo lordo 348,00 EUR", "A1"),
    V("premio", "Premio annuo", "384,00 EUR annui", "Premio annuo lordo 348,00 EUR", "A1"),        # numero sbagliato
    V("capitale", "Capitale caso morte", "200.000,00 EUR costante",
      "Capitale assicurato caso morte 200.000,00 EUR - costante per tutta la durata", "A1"),
    V("durata", "Durata", "15 anni", "Durata del contratto 15 anni", "A1", deroga=True),            # deroga a torto
    V("frazionamento", "Frazionamento", "annuale, semestrale, trimestrale, mensile senza maggiorazione",
      "annuale, semestrale, trimestrale, mensile - senza maggiorazione", "A1"),
    V("carenza", "Carenza per malattia", "6 mesi",
      "Per i decessi dovuti a malattia la garanzia e' operante trascorsi 6 mesi dalla decorrenza", "A3"),  # pagina sbagliata
    V("carenza", "Carenza per malattia ridotta", "3 mesi con referto ematochimico",
      "il periodo di carenza per decesso da malattia e' ridotto a 3 mesi a fronte della consegna di "
      "referto ematochimico", "A4", deroga=True),
    V("esclusione", "Immersione subacquea", "esclusa oltre 40 metri",
      "Immersione subacquea con autorespiratore oltre i 40 metri di profondita'", "A2"),
    V("esclusione", "Alpinismo", "esclusa roccia oltre III grado e ascensioni oltre 4.000 metri",
      "Alpinismo con scalata di roccia oltre il III grado e ascensioni oltre 4.000 metri", "A2"),
    V("esclusione", "Alpinismo su ghiacciaio e immersione entro 40 metri",
      "non escluse (pratica non agonistica, ghiacciaio con guida)",
      "In parziale deroga all'Art. 3, la pratica non agonistica di immersione subacquea entro i 40 metri "
      "e dell'alpinismo su ghiacciaio con guida abilitata non costituisce causa di esclusione", "A4", deroga=True),
    V("esclusione", "Epidemie e pandemie", "escluse",
      "sono escluse le epidemie e pandemie dichiarate dall'autorita'", "A2"),                            # inventata
    V("garanzia_complementare", "Invalidita' permanente", "da infortunio o malattia, dal 60%, compresa",
      "Invalidita' permanente da infortunio o malattia compresa nel premio base anticipazione del capitale "
      "assicurato se il grado di invalidita' riconosciuto e' pari o superiore al 60%", "A3", inclusa=True),
    V("garanzia_complementare", "Doppio capitale per infortunio", "compreso, maggiorazione massima 100.000 EUR",
      "raddoppio del capitale, con il limite massimo di 100.000 EUR di maggiorazione", "A3", inclusa=True),
    V("garanzia_complementare", "Inabilita' temporanea totale",
      "opzionale non attivata: 40 EUR al giorno, franchigia 30 giorni, massimo 365 gg, 72 EUR/anno",
      "opzionale, non attivata in proposta indennita' giornaliera di 40 EUR, franchigia di 30 giorni, "
      "massimo 365 gg 72 EUR/anno", "A3", inclusa=False),
    V("garanzia_complementare", "Malattie gravi", "non prevista", "Malattie gravi non prevista dal prodotto",
      "A3", inclusa=False),
    V("servizi", "Assistenza", "centrale operativa 24 ore su 24, medico urgente, rimpatrio fino a 15.000 EUR",
      "centrale operativa attiva 24 ore su 24 per consulenza medica telefonica", "A3"),
    V("costi", "Caricamenti", "22,0% del premio lordo", "Caricamenti sul premio 22,0% del premio lordo", "A3"),
    V("recesso", "Recesso", "30 giorni, rimborso meno 25 EUR di spese",
      "premio versato al netto delle spese di emissione, pari a 25 EUR", "A3"),
    V("sconti", "Sconto gia' clienti", "5% per chi ha gia' un'altra polizza, non applicato",
      "e' riconosciuto uno sconto commerciale del 5% sul premio annuo per i contraenti gia' titolari di altra "
      "polizza della Compagnia", "A4", deroga=True),
    V("non_contestabilita", "Non contestabilita'", "dopo 24 mesi",
      "Trascorsi 24 mesi dalla decorrenza, la Compagnia non puo' contestare le dichiarazioni", "A2"),
    V("riscatto", "Riscatto", "non previsto", "Il contratto non prevede valore di riscatto ne' di riduzione", "A1"),
]}
# di proposito manca la «prestazione» di A: deve uscire «non l'ho trovato, ho cercato in ...»
# (il «30 giorni» del recesso non e' nella citazione: la voce deve essere scartata)

LETTURA_B = {"ramo": "vita - temporanea caso morte", "compagnia": "NORDEST VITA S.p.A.",
             "prodotto": "NE SERENA CASA E FAMIGLIA", "voci": [
    V("prodotto", "Tipo di contratto", "temporanea caso morte a capitale decrescente",
      "Polizza temporanea caso morte a capitale decrescente", "B1"),
    V("premio", "Premio annuo", "156,00 EUR", "Premio annuo 156,00 EUR", "B1"),
    V("capitale", "Capitale iniziale", "200.000,00 EUR iniziale", "Capitale iniziale assicurato 200.000,00 EUR", "B1"),
    V("andamento_capitale", "Andamento del capitale", "decrescente fino a zero, capitale medio 100.000 EUR",
      "decrescente in quote annue costanti fino a zero al termine del 15. anno (capitale medio 100.000 EUR circa)",
      "B1"),
    V("durata", "Durata", "15 anni", "Durata 15 anni", "B1"),
    V("prestazione", "Prestazione", "capitale residuo alla data del decesso",
      "la Societa' liquida ai beneficiari il capitale residuo alla data del decesso", "B2"),
    V("carenza", "Carenza per malattia", "12 mesi",
      "Per i decessi conseguenti a malattia la garanzia opera decorsi 12 mesi dalla data di effetto", "B2"),
    V("esclusione", "Immersione subacquea", "esclusa oltre 30 metri", "Immersione subacquea oltre i 30 metri di profondita'", "B2"),
    V("esclusione", "Alpinismo", "escluso senza limite di quota o grado",
      "Alpinismo, arrampicata, speleologia senza limite di quota o grado", "B2"),
    V("esclusione", "Epidemie e pandemie", "escluse se dichiarato lo stato di emergenza",
      "Epidemie e pandemie dichiarate quando l'autorita' sanitaria abbia dichiarato lo stato di emergenza", "B2"),
    V("garanzia_complementare", "Invalidita' permanente", "solo da infortunio, dal 66%",
      "residuo se l'invalidita' riconosciuta e' pari o superiore al 66%; esclusa l'invalidita' da malattia",
      "B3", inclusa=True),
    V("garanzia_complementare", "Doppio capitale per infortunio", "non prevista",
      "Doppio capitale per infortunio non prevista", "B3", inclusa=False),
    V("garanzia_complementare", "Inabilita' temporanea totale",
      "attivata: 30 EUR al giorno, franchigia 60 giorni, massimo 180 gg, 48 EUR compresi nei 156 EUR",
      "indennita' giornaliera di 30 EUR, franchigia di 60 giorni, massimo 180 gg 48 EUR/anno compresi nei 156 EUR",
      "B3", inclusa=True),
    V("garanzia_complementare", "Malattie gravi", "anticipo del 30% alla diagnosi di 4 patologie",
      "anticipazione del 30% del capitale residuo alla diagnosi di una delle 4 patologie", "B3", inclusa=True),
    V("costi", "Caricamenti", "31,5% del premio lordo", "Caricamenti sul premio 31,5% del premio lordo", "B3"),
    V("revisione_premio", "Revisione della tariffa", "al 5. e al 10. anno",
      "la Societa' si riserva la revisione della tariffa al 5. e al 10. anno", "B3"),
]}


def _id(voci, chiave, parola):
    return next(v["id"] for v in voci if v["chiave"] == chiave and parola.lower() in v["nome"].lower())


def giudizio_finto(dati):
    A, B = dati["A"], dati["B"]
    return {
        "righe": [
            {"voce": "Premio", "a": [_id(A, "premio", "premio")], "b": [_id(B, "premio", "premio")],
             "giudizio": "diverso", "perche": "348,00 EUR contro 156,00 EUR: 192 EUR di differenza, su capitali diversi."},
            {"voce": "Carenza", "a": [_id(A, "carenza", "malattia")], "b": [_id(B, "carenza", "malattia")],
             "giudizio": "meglio_a", "perche": "6 mesi contro 12 mesi."},
            {"voce": "Immersione", "a": [_id(A, "esclusione", "immersione")], "b": [_id(B, "esclusione", "immersione")],
             "giudizio": "meglio_a", "perche": "A esclude oltre 40 metri, B gia' oltre 30 metri."},
            {"voce": "Doppio capitale", "a": [_id(A, "garanzia_complementare", "doppio")],
             "b": [_id(B, "garanzia_complementare", "doppio")], "giudizio": "solo_a", "perche": ""},
            # incoerente: A le malattie gravi NON le ha. Il codice deve correggere il giudizio
            {"voce": "Malattie gravi", "a": [_id(A, "garanzia_complementare", "malattie")],
             "b": [_id(B, "garanzia_complementare", "malattie")], "giudizio": "solo_a", "perche": ""},
            {"voce": "Fantasma", "a": ["A-999"], "b": ["B-999"], "giudizio": "meglio_a", "perche": "non esiste"},
        ],
        "trappole": [
            {"testo": "Il premio di B (156,00 EUR) e' su capitale decrescente: per 1.000 EUR di capitale medio "
                      "A costa 1,74 e B 1,56.", "voci": [_id(A, "premio", "premio"), _id(B, "premio", "premio")]},
            {"testo": "Con B il cliente risparmia 250 EUR l'anno.", "voci": []},                 # numero inventato
        ],
        "azioni": [{"testo": "Chiedere al cliente il referto ematochimico: con A la carenza scende a 3 mesi.",
                    "voci": [_id(A, "carenza", "ridotta")]}],
        "sintesi": "A costa di piu' ma copre un capitale costante e piu' rischi.",
    }


class ModelloFinto:
    def __init__(self):
        self.chiamate = []

    def __call__(self, sistema, utente, max_tokens, temperatura):
        self.chiamate.append(utente)
        if utente.startswith("FASCICOLO:"):
            return json.dumps(LETTURA_A if "[[A1]]" in utente else LETTURA_B)
        return "```json\n" + json.dumps(giudizio_finto(json.loads(utente))) + "\n```"


# ------------------------------------------------------------------ preparazione
@pytest.fixture(scope="module")
def pdf(tmp_path_factory):
    return polizze_prova.genera(str(tmp_path_factory.mktemp("polizze")))


def _doc(path, titolo, did):
    chunks, _ = docs.extract(path, ".pdf")
    return {"id": did, "titolo": titolo, "chunks": chunks}


@pytest.fixture
def modello(monkeypatch):
    m = ModelloFinto()
    monkeypatch.setattr(cp, "scrivi", m)
    return m


@pytest.fixture
def risultato(pdf, modello):
    return cp.confronta([_doc(pdf[0], "Polizza Versilia", "a1")], [_doc(pdf[1], "Polizza Nordest", "b1")])


def _voce(lato, chiave, parola):
    return next(v for v in lato["voci"] if v["chiave"] == chiave and parola.lower() in v["nome"].lower())


# ------------------------------------------------------------------ le forme del testo
def test_citazione_regge_le_forme_dei_pdf():
    # tabella con una cella per riga e celle vicine in mezzo
    testo = ("Invalidita' permanente da\ninfortunio o malattia\ncompresa nel premio\nbase\n"
             "anticipazione del capitale\nassicurato se il grado di\ninvalidita' riconosciuto e'\n"
             "pari o superiore al 60%\ninclusa")
    assert cp.trova_citazione("Invalidità permanente da infortunio o malattia compresa nel premio base "
                              "anticipazione del capitale assicurato se il grado di invalidità riconosciuto è "
                              "pari o superiore al 60%", testo)
    assert cp.trova_citazione("Invalidita' permanente da infortunio, pari o superiore al 60%", testo)
    # parola spezzata a fine riga, legatura, apostrofo tipografico
    assert cp.trova_citazione("il capitale assicurato è definito", "il capitale assicu-\nrato e' definito")
    assert cp.trova_citazione("la garanzia è efficace", "la garanzia e’ eﬃcace".replace("eﬃ", "eﬃ"))
    assert cp.trova_citazione("franchigia di 30 giorni", "franchigia di\n30 giorni")
    assert cp.trova_citazione("capitale di 200.000 EUR", "capitale di 200000 EUR")
    # NON passano: parole in un altro ordine, numero diverso, parole troppo lontane, troppo corta
    assert not cp.trova_citazione("superiore al 66%", "pari o superiore al 60%")
    assert not cp.trova_citazione("malattia o infortunio da permanente", testo)
    lontano = "carenza " + "parola " * 80 + "di sei mesi"
    assert not cp.trova_citazione("carenza di sei mesi", lontano)
    assert not cp.trova_citazione("60%", testo)


def test_numeri_in_forma_unica():
    assert cp.numeri("200.000,00 EUR") == {"200000"}
    assert cp.numeri("22,0% e 31,5%") == {"22", "31.5"}
    assert cp.numeri("tasso 0,8 g/l") == {"0.8"}
    assert cp.numeri("massimo 1.000 euro e 4.000 metri") == {"1000", "4000"}


# ------------------------------------------------------------------ verifica della lettura
def test_lettura_verificata_sulle_polizze_vere(risultato):
    a, b = risultato["a"], risultato["b"]
    scartate = risultato["controlli"]["scartate"]
    motivi = {(s.get("nome"), s.get("motivo", "").split(":")[0]) for s in scartate}
    # inventata: fuori; numero sbagliato: fuori; numero non nella citazione: fuori
    assert ("Epidemie e pandemie", "citazione non trovata nel fascicolo") in motivi
    assert ("Premio annuo", "numeri non presenti nella citazione") in motivi
    assert ("Recesso", "numeri non presenti nella citazione") in motivi
    assert not any(v["nome"] == "Epidemie e pandemie" for v in a["voci"])
    assert [v["valore"] for v in a["voci"] if v["chiave"] == "premio"] == ["348,00 EUR annui"]
    # pagina sbagliata: corretta, non scartata
    assert _voce(a, "carenza", "per malattia")["label"] == "pagina 2"
    # deroga dichiarata a torto su una pagina senza deroga: non e' una deroga
    assert _voce(a, "durata", "durata")["deroga"] is False
    # le celle di tabella spezzate su piu' righe passano
    assert _voce(a, "garanzia_complementare", "invalidita")["valore"].endswith("compresa")
    assert _voce(b, "garanzia_complementare", "inabilita")["inclusa"] is True
    assert a["pagine_deroga"] == ["A4"] and b["pagine_deroga"] == []


def test_la_deroga_vale_solo_per_quello_che_dice(risultato):
    a = risultato["a"]
    base = _voce(a, "carenza", "per malattia")
    ridotta = _voce(a, "carenza", "ridotta")
    # la carenza ridotta vale solo «a fronte del referto»: i 6 mesi restano, modificati
    assert ridotta["deroga"] and ridotta["condizionata"]
    assert base["superata"] is False and base["modificata_da"] == [ridotta["id"]]
    # «in parziale deroga» sulle esclusioni: le esclusioni base restano, marcate
    deroga_escl = _voce(a, "esclusione", "ghiacciaio")
    assert deroga_escl["id"] in _voce(a, "esclusione", "immersione")["modificata_da"]
    assert deroga_escl["id"] in next(v for v in a["voci"] if v["nome"] == "Alpinismo")["modificata_da"]
    assert _voce(a, "esclusione", "immersione")["superata"] is False
    assert _voce(a, "sconti", "sconto")["condizionata"] is True


def test_deroga_senza_condizioni_sostituisce():
    voci = [{"id": "A-1", "chiave": "carenza", "nome": "Carenza", "valore": "6 mesi", "deroga": False,
             "citazione": "carenza di 6 mesi"},
            {"id": "A-2", "chiave": "carenza", "nome": "Carenza", "valore": "3 mesi", "deroga": True,
             "citazione": "in deroga all'art. 2 la carenza e' di 3 mesi"}]
    cp._applica_deroghe(voci)
    assert voci[0]["superata"] is True and voci[0]["modificata_da"] == ["A-2"]


def test_premio_sulla_stessa_base(risultato):
    ba, bb = risultato["a"]["base_premio"], risultato["b"]["base_premio"]
    assert ba["calcolabile"] and ba["andamento"] == "costante" and ba["per_mille"] == 1.74
    assert bb["calcolabile"] and bb["andamento"] == "decrescente" and bb["capitale_medio"] == 100000
    assert bb["per_mille"] == 1.56 and "scritto" in bb["come"]


def test_base_premio_non_si_inventa():
    v = lambda chiave, valore, cit: {"chiave": chiave, "valore": valore, "citazione": cit, "superata": False}
    voci = [v("premio", "156,00 EUR", "premio 156,00 EUR"), v("capitale", "200.000 EUR", "capitale 200.000 EUR"),
            v("andamento_capitale", "decrescente", "capitale decrescente secondo il piano allegato")]
    out = cp.base_premio(voci)
    assert out["calcolabile"] is False and "piano di decrescenza" in out["motivo"]
    lineare = voci[:2] + [v("andamento_capitale", "decrescente", "decrescente in quote annue costanti fino a zero")]
    assert cp.base_premio(lineare)["capitale_medio"] == 100000
    assert cp.base_premio(voci[1:])["motivo"] == "premio non trovato"


# ------------------------------------------------------------------ verifica del giudizio
def test_giudizio_controllato(risultato):
    righe = risultato["righe"]
    per_voce = {r["voce"]: r for r in righe}
    assert "Fantasma" not in per_voce
    assert per_voce["Doppio capitale"]["giudizio"] == "solo_a"
    assert per_voce["Malattie gravi"]["giudizio"] == "diverso"           # corretto dal codice
    assert "192 EUR" in per_voce["Premio"]["perche"]                     # differenza: ammessa
    trappole = [t["testo"] for t in risultato["trappole"]]
    assert len(trappole) == 1 and "1,74" in trappole[0]
    assert not any("250" in t for t in trappole)
    assert any(s.get("dove") == "trappole" and "250" in s["motivo"] for s in risultato["controlli"]["scartate"])
    assert risultato["azioni"][0]["testo"].startswith("Chiedere al cliente il referto")


def test_nessuna_voce_verificata_sparisce(risultato):
    nelle_righe = {i for r in risultato["righe"] for i in r["a"] + r["b"]}
    for lato in (risultato["a"], risultato["b"]):
        for v in lato["voci"]:
            if not v["superata"]:
                assert v["id"] in nelle_righe, v
    assert any(r["giudizio"] == "non_confrontata" for r in risultato["righe"])
    # la stessa voce di A e di B, non confrontata dal modello, sta sulla stessa riga
    durata = [r for r in risultato["righe"] if r["voce"] == "Durata"]
    assert len(durata) == 1 and durata[0]["a"] and durata[0]["b"]
    invalidita = [r for r in risultato["righe"] if r["voce"].startswith("Invalidita")]
    assert len(invalidita) == 1 and invalidita[0]["a"] and invalidita[0]["b"]


def test_voce_essenziale_mancante_dichiarata(risultato):
    mancanti = {n["chiave"]: n["testo"] for n in risultato["a"]["non_trovate"]}
    assert "prestazione" in mancanti
    assert mancanti["prestazione"].startswith("Non l'ho trovato. Ho cercato in: Polizza Versilia pagina 1")
    assert "prestazione" not in {n["chiave"] for n in risultato["b"]["non_trovate"]}


def test_sintesi_con_numero_inventato_tolta(pdf, monkeypatch):
    m = ModelloFinto()

    def con_sintesi_falsa(s, u, t, temp):
        out = m(s, u, t, temp)
        if not u.startswith("FASCICOLO:"):
            dati = json.loads(out.strip("`").replace("json\n", "", 1))
            dati["sintesi"] = "B fa risparmiare il 40% a parita' di tutto."
            return json.dumps(dati)
        return out
    monkeypatch.setattr(cp, "scrivi", con_sintesi_falsa)
    r = cp.confronta([_doc(pdf[0], "A", "a1")], [_doc(pdf[1], "B", "b1")])
    assert r["sintesi"] == ""
    assert any(s.get("dove") == "sintesi" for s in r["controlli"]["scartate"])


def test_fascicolo_lungo_letto_a_lotti(monkeypatch):
    monkeypatch.setattr(cp, "MAX_CARATTERI_LOTTO", 500)
    letti = []

    def finto(s, u, t, temp):
        letti.append(u)
        return json.dumps({"voci": []})
    monkeypatch.setattr(cp, "scrivi", finto)
    pagine = [{"label": f"pagina {n}", "text": "testo " * 60} for n in range(1, 6)]
    cp.leggi_lato(cp.fascicolo("A", [{"id": "x", "titolo": "Lungo", "chunks": pagine}]))
    assert len(letti) == 5
    assert all(u.count("[[A") == 1 for u in letti)


def test_pagine_senza_testo_e_parti_riunite():
    chunks = [{"label": "pagina 1", "text": "uno"}, {"label": "pagina 3 · parte 1", "text": "tre a"},
              {"label": "pagina 3 · parte 2", "text": "tre b"}]
    f = cp.fascicolo("B", [{"id": "x", "titolo": "Scansione", "chunks": chunks}])
    assert [p["label"] for p in f["pagine"]] == ["pagina 1", "pagina 3"]
    assert f["pagine"][1]["testo"] == "tre a\ntre b"
    assert f["senza_testo"] == ["«Scansione» pagina 2"]
    with pytest.raises(HTTPException) as e:
        cp.confronta([{"id": "a", "titolo": "A", "chunks": []}], [{"id": "b", "titolo": "B", "chunks": chunks}])
    assert e.value.status_code == 422 and "scansione" in e.value.detail


def test_json_rotto_riprovato(pdf, monkeypatch):
    m = ModelloFinto()
    rotti = {"n": 0}

    def una_volta_rotto(s, u, t, temp):
        if u.startswith("FASCICOLO:") and "[[B1]]" in u and not rotti["n"]:
            rotti["n"] += 1
            return "ecco le voci: {non json"
        return m(s, u, t, temp)
    monkeypatch.setattr(cp, "scrivi", una_volta_rotto)
    r = cp.confronta([_doc(pdf[0], "A", "a1")], [_doc(pdf[1], "B", "b1")])
    assert rotti["n"] == 1 and r["b"]["voci"]


# ------------------------------------------------------------------ Excel
def test_excel(risultato):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(cp.excel(risultato)))
    assert wb.sheetnames == ["Confronto", "Base premio", "Trappole e azioni", "Dettaglio A", "Dettaglio B", "Controlli"]
    ws = wb["Confronto"]
    testi = [str(c.value or "") for r in ws.iter_rows() for c in r]
    assert any("[DEROGA, a condizione]" in t for t in testi)
    assert any("modificato dalla deroga" in t for t in testi)
    assert any(t.startswith("Non l'ho trovato") for t in testi)
    assert wb["Base premio"]["B7"].value.startswith("=IF(")
    assert wb["Base premio"]["B4"].value == 348 and wb["Base premio"]["C6"].value == 100000
    assert wb["Controlli"].max_row >= 4 + len(risultato["controlli"]["scartate"])
    assert ws["A1"].font.name == "Arial"


# ------------------------------------------------------------------ rotte: chiavi, archivi, credito
KEYS = {"agenzia": "chiave-agenzia", "altra": "chiave-altra"}


@pytest.fixture
def http(pdf, modello, monkeypatch, tmp_path):
    monkeypatch.setenv("DOCUMENT_UPLOAD_ROOT", str(tmp_path / "archivio"))
    monkeypatch.setenv("APPLAUDI_TENANT_CLIENT_MAP", json.dumps({"agenzia": "agenzia", "altra": "altra"}))

    def verifica(tenant, chiave):
        if KEYS.get(tenant) != chiave:
            raise HTTPException(401, "Chiave non valida")
        return tenant
    monkeypatch.setattr(docs, "tenant_verifier", verifica)
    app = FastAPI()
    app.include_router(docs.router)
    app.include_router(cp.router)
    client = TestClient(app)
    ids = {}
    for tenant, path in (("agenzia", pdf[0]), ("agenzia", pdf[1]), ("altra", pdf[1])):
        with open(path, "rb") as f:
            r = client.post(f"/api/documents/{tenant}/upload", files={"file": (path.split("/")[-1], f, "application/pdf")},
                            headers={"X-Applaudi-Tenant-Key": KEYS[tenant]})
        assert r.status_code == 200, r.text
        assert r.json()["document"]["tipo"] == "polizza"      # non «contratto», non «promozione»
        ids.setdefault(tenant, []).append(r.json()["document"]["id"])
    return client, ids


def test_rotta_confronta_e_scarica(http):
    client, ids = http
    h = {"X-Applaudi-Tenant-Key": KEYS["agenzia"]}
    r = client.post("/api/confronta/agenzia", json={"a": [ids["agenzia"][0]], "b": [ids["agenzia"][1]],
                                                    "nome_a": "La nostra", "nome_b": "Concorrente"}, headers=h)
    assert r.status_code == 200, r.text
    dati = r.json()
    assert dati["a"]["nome"] == "La nostra" and dati["b"]["nome"] == "Concorrente"
    assert client.get("/api/confronta/agenzia", headers=h).json()["confronti"][0]["id"] == dati["id"]
    x = client.get(f"/api/confronta/agenzia/{dati['id']}/excel", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    # un'altra agenzia non vede il confronto, ne' con il suo id ne' con la sua chiave
    altra = {"X-Applaudi-Tenant-Key": KEYS["altra"]}
    assert client.get(f"/api/confronta/altra/{dati['id']}", headers=altra).status_code == 404
    assert client.get(f"/api/confronta/agenzia/{dati['id']}", headers=altra).status_code == 401
    assert client.get("/api/confronta/agenzia/../../x", headers=h).status_code == 404


def test_rotta_rifiuta(http, monkeypatch):
    client, ids = http
    h = {"X-Applaudi-Tenant-Key": KEYS["agenzia"]}
    a, b = ids["agenzia"]
    # il documento di un'altra agenzia non esiste qui
    assert client.post("/api/confronta/agenzia", json={"a": [a], "b": [ids["altra"][0]]}, headers=h).status_code == 404
    assert client.post("/api/confronta/agenzia", json={"a": [a], "b": [a]}, headers=h).status_code == 400
    assert client.post("/api/confronta/agenzia", json={"a": [], "b": [b]}, headers=h).status_code == 400
    assert client.post("/api/confronta/agenzia", json={"a": [a], "b": [b]}, headers={}).status_code == 401
    monkeypatch.setattr(cp, "credito_ok", lambda cliente: False)
    assert client.post("/api/confronta/agenzia", json={"a": [a], "b": [b]}, headers=h).status_code == 402


def test_i_token_vanno_al_cliente_giusto(http, monkeypatch):
    client, ids = http
    visti = []
    m = ModelloFinto()

    def contato(s, u, t, temp):
        visti.append(docs.consumo_cliente.get())
        return m(s, u, t, temp)
    monkeypatch.setattr(cp, "scrivi", contato)
    r = client.post("/api/confronta/agenzia", json={"a": [ids["agenzia"][0]], "b": [ids["agenzia"][1]]},
                    headers={"X-Applaudi-Tenant-Key": KEYS["agenzia"]})
    assert r.status_code == 200
    assert visti and set(visti) == {"agenzia"}          # anche nelle letture in parallelo
