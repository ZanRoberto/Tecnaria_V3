"""Collaudo dei documenti aziendali sul CORE VERO (app.py + indice LAGO reale).

Simulato SOLO il modello linguistico. Rotte, chiavi, archivio su disco, estrazione,
scheda della fonte, validita' nel tempo, cambio cliente, registro: codice vero.

Cosa deve restare vero:
- il catalogo LAGO non cambia: stesse pagine, stessi codici, stesso messaggio al modello;
- un documento non confermato non entra in nessuna risposta;
- una promozione vale solo nel suo periodo, e la memoria delle risposte se ne accorge;
- un documento sostituito esce dalle risposte ma resta nello storico;
- un documento riservato non esce mai su WhatsApp;
- un'azienda non vede mai i documenti di un'altra; l'azienda la decide il server.

Esecuzione:  python -m pytest -q test_documenti_aziendali.py
"""
import asyncio
import datetime as dt
import io
import json
import types

import pytest
from fastapi.testclient import TestClient

from test_whatsapp_core import FakeLLM, core  # noqa: F401  (fixture condivisa)

KEYS = {'lago': 'chiave-lago', 'covolo': 'chiave-covolo', 'roberto': 'chiave-roberto',
        'estraneo': 'chiave-estraneo'}
DOMANDA = 'Qual e la procedura di collaudo antiribaltamento per il mobile TV?'
FRASE_LAGO = 'Procedura di collaudo antiribaltamento: fissare il mobile TV a muro con 4 tasselli ZX-77.'
FRASE_COVOLO = 'Procedura di collaudo antiribaltamento COVOLO: usare la staffa segreta CVZ-SEGRETA.'
PROMO = ('PROMOZIONE AUTUNNO. Mobile TV N.O.W. codice 211605 a prezzo promozionale 2.500 euro '
         'invece del listino. Valida dal 01/10/2026 fino al 30/11/2026. Solo finitura vetro opaco.')
DOMANDA_PROMO = 'Che prezzo ha il mobile TV 211605? ci sono promozioni?'


@pytest.fixture
def env(core, monkeypatch, tmp_path):  # noqa: F811
    import plaud_remote
    import documenti_aziendali as docs
    monkeypatch.setenv('DOCUMENT_UPLOAD_ROOT', str(tmp_path / 'archivio'))
    monkeypatch.delenv('APPLAUDI_TENANT_CLIENT_MAP', raising=False)
    monkeypatch.setattr(plaud_remote, 'TENANT_KEYS', dict(KEYS))
    monkeypatch.setattr(docs, 'oggi', lambda: dt.date(2026, 10, 17))
    llm = FakeLLM(core, answer='Risposta di prova.')
    monkeypatch.setattr(core, 'client', llm)
    monkeypatch.setattr(core, 'openai_client', None)
    monkeypatch.setitem(core.COLLAUDO_STATE, 'running', None)
    for cliente in ('lago', 'covolo'):          # archivio nuovo per ogni test
        core._documenti_cambiati(cliente)
    for stato in core.STATO_PER_CLIENTE.values():
        stato['ANSWER_CACHE'].clear()
    core.ANSWER_CACHE.clear()
    return type('E', (), {'core': core, 'docs': docs, 'llm': llm,
                          'http': TestClient(core.app), 'root': tmp_path / 'archivio'})


def H(tenant):
    return {'X-Applaudi-Tenant-Key': KEYS[tenant]}


def upload(env, tenant, name, data, key=None):
    return env.http.post(f'/api/documents/{tenant}/upload',
                         files={'file': (name, io.BytesIO(data), 'application/octet-stream')},
                         headers={'X-Applaudi-Tenant-Key': key or KEYS[tenant]})


def scheda(env, tenant, doc_id, **campi):
    return env.http.post(f'/api/documents/{tenant}/{doc_id}/scheda', json=campi, headers=H(tenant))


def carica(env, tenant, name, data, **campi):
    """Carica e conferma (il caso normale): restituisce il documento confermato."""
    r = upload(env, tenant, name, data)
    assert r.status_code == 200, r.text
    r = scheda(env, tenant, r.json()['document']['id'], **campi)
    assert r.status_code == 200, r.text
    return r.json()['document']


def ask(env, cliente, question=DOMANDA):
    env.llm.calls.clear()
    r = env.http.post('/api/ask', json={'question': '/catalogo ' + question, 'cliente': cliente})
    assert r.status_code == 200, r.text
    prompts = [c['text'] for c in env.llm.calls if 'RICHIESTA ORIGINALE' in c['text']]
    return prompts[0] if prompts else ''


def without_block(prompt):
    head, sep, tail = prompt.partition('\n\nDOCUMENTI AZIENDALI CARICATI')
    if not sep:
        return prompt
    return head + tail[tail.index('\n\nFormula ora'):]


# ---------------------------------------------------------------- il catalogo non cambia

def test_catalogo_lago_identico_con_e_senza_documenti(env):
    core = env.core
    core.attiva_cliente('lago')
    pagine, codici = len(core.DOCUMENT_PAGES), len(core.CODE_ROWS)
    prima = ask(env, 'lago')
    assert prima and 'DOCUMENTI AZIENDALI CARICATI' not in prima
    carica(env, 'lago', 'collaudo.txt', FRASE_LAGO.encode(), tipo='manuale')
    dopo = ask(env, 'lago')
    core.attiva_cliente('lago')
    assert len(core.DOCUMENT_PAGES) == pagine and len(core.CODE_ROWS) == codici
    assert 'ZX-77' in dopo and '«collaudo»' in dopo
    assert without_block(dopo) == prima          # stesso catalogo, blocco in piu' in coda


def test_documento_non_pertinente_non_entra(env):
    carica(env, 'lago', 'ferie.txt', b'Calendario ferie estive del magazzino di Bassano.')
    prompt = ask(env, 'lago', 'Avete un mobile TV largo circa 160 cm?')
    assert 'DOCUMENTI AZIENDALI CARICATI' not in prompt


# ---------------------------------------------------------------- scheda e conferma

def test_non_confermato_non_entra_nelle_risposte(env):
    r = upload(env, 'lago', 'collaudo.txt', FRASE_LAGO.encode())
    doc = r.json()['document']
    assert doc['stato'] == 'da_confermare' and doc['situazione'] == 'da_confermare'
    assert 'ZX-77' not in ask(env, 'lago')
    assert scheda(env, 'lago', doc['id'], tipo='manuale').status_code == 200
    assert 'ZX-77' in ask(env, 'lago')


def test_proposte_lette_dal_documento(env):
    r = upload(env, 'lago', 'Promo Autunno.txt', PROMO.encode())
    doc = r.json()['document']
    p = doc['proposte']
    assert p['tipo'] == 'promozione' and p['titolo'] == 'Promo Autunno'
    assert p['valido_dal'] == '2026-10-01' and p['valido_fino'] == '2026-11-30'
    assert doc['valido_fino'] is None              # proposta, non ancora decisa
    assert '2.500' not in ask(env, 'lago', DOMANDA_PROMO)


def test_promozione_senza_scadenza_rifiutata(env):
    doc = upload(env, 'lago', 'promo.txt', PROMO.encode()).json()['document']
    r = scheda(env, 'lago', doc['id'], tipo='promozione', valido_dal='2026-10-01')
    assert r.status_code == 400 and 'fine validit' in r.json()['detail']
    r = scheda(env, 'lago', doc['id'], tipo='promozione', valido_dal='2026-12-01', valido_fino='2026-11-30')
    assert r.status_code == 400


def test_stesso_file_due_volte(env):
    assert upload(env, 'lago', 'a.txt', FRASE_LAGO.encode()).status_code == 200
    r = upload(env, 'lago', 'copia.txt', FRASE_LAGO.encode())
    assert r.status_code == 409 and 'gia' in r.json()['detail']


# ---------------------------------------------------------------- la promozione nasce e muore

@pytest.mark.parametrize('giorno,atteso', [
    (dt.date(2026, 9, 20), 'in_arrivo'),
    (dt.date(2026, 10, 1), 'vigente'),
    (dt.date(2026, 11, 30), 'vigente'),
    (dt.date(2026, 12, 1), 'scaduto'),
])
def test_ciclo_di_vita_promozione(env, monkeypatch, giorno, atteso):
    carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione', titolo='Promo Autunno',
           valido_dal='2026-10-01', valido_fino='2026-11-30')
    monkeypatch.setattr(env.docs, 'oggi', lambda: giorno)
    elenco = env.http.get('/api/documents/lago', headers=H('lago')).json()['documents']
    assert elenco[0]['situazione'] == atteso
    prompt = ask(env, 'lago', DOMANDA_PROMO)
    if atteso == 'scaduto':
        assert '2.500' not in prompt and 'Promo Autunno' not in prompt
    elif atteso == 'in_arrivo':
        assert 'NON ANCORA IN VIGORE' in prompt and '01/10/2026' in prompt
    else:
        assert '2.500' in prompt and 'NON ANCORA IN VIGORE' not in prompt
        assert 'valido dal 01/10/2026 al 30/11/2026' in prompt


def test_memoria_delle_risposte_sa_che_la_promo_e_morta(env, monkeypatch):
    llm = FakeLLM(env.core, answer='RISPOSTA\nListino pag. 60.')
    monkeypatch.setattr(env.core, 'client', llm)
    env.llm = llm
    carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione',
           valido_dal='2026-10-01', valido_fino='2026-11-30')
    assert '2.500' in ask(env, 'lago', DOMANDA_PROMO)          # durante: calcolata e memorizzata
    monkeypatch.setattr(env.docs, 'oggi', lambda: dt.date(2026, 12, 1))
    dopo = ask(env, 'lago', DOMANDA_PROMO)                      # dopo: NON dalla memoria
    assert dopo and '2.500' not in dopo


# ---------------------------------------------------------------- sostituzione e storico

def test_sostituzione_esce_dalle_risposte_resta_nello_storico(env):
    vecchia = carica(env, 'lago', 'promo1.txt', PROMO.encode(), tipo='promozione',
                     titolo='Promo Autunno', valido_dal='2026-10-01', valido_fino='2026-11-30')
    nuova_testo = PROMO.replace('2.500', '2.400').replace('AUTUNNO', 'AUTUNNO BIS')
    nuova = upload(env, 'lago', 'promo2.txt', nuova_testo.encode()).json()['document']
    r = scheda(env, 'lago', nuova['id'], tipo='promozione', titolo='Promo Autunno bis',
               valido_dal='2026-10-01', valido_fino='2026-11-30', sostituisce=vecchia['id'])
    assert r.status_code == 200
    prompt = ask(env, 'lago', DOMANDA_PROMO)
    assert '2.400' in prompt and '2.500' not in prompt
    attivi = env.http.get('/api/documents/lago', headers=H('lago')).json()['documents']
    assert [d['titolo'] for d in attivi] == ['Promo Autunno bis']
    storico = env.http.get('/api/documents/lago?storico=true', headers=H('lago')).json()['documents']
    old = next(d for d in storico if d['id'] == vecchia['id'])
    assert old['stato'] == 'sostituito' and old['sostituito_da'] == nuova['id']
    assert any('sostituito' in s['cosa'] for s in old['storia'])
    assert scheda(env, 'lago', old['id'], titolo='x').status_code == 409   # lo storico non si riscrive


def test_modifica_scheda_resta_nella_storia_e_cambia_versione(env):
    doc = carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione',
                 valido_dal='2026-10-01', valido_fino='2026-11-30')
    d = scheda(env, 'lago', doc['id'], valido_fino='2026-12-31').json()['document']   # proroga
    assert d['versione'] == doc['versione'] + 1
    assert d['storia'][-1]['cambi'] == {'valido_fino': '2026-12-31'}


def test_archiviato_non_entra(env):
    doc = carica(env, 'lago', 'collaudo.txt', FRASE_LAGO.encode(), tipo='manuale')
    assert 'ZX-77' in ask(env, 'lago')
    assert scheda(env, 'lago', doc['id'], stato='archiviato').status_code == 200
    assert 'ZX-77' not in ask(env, 'lago')


def test_mail_eml_dati_certi_e_durata(env, monkeypatch):
    eml = (b'From: Mario Rossi <mario@esempio.it>\r\nTo: vendite@esempio.it\r\n'
           b'Date: Mon, 12 Oct 2026 10:00:00 +0200\r\nSubject: Richiesta collaudo antiribaltamento\r\n'
           b'Content-Type: text/plain; charset=utf-8\r\n\r\n'
           b'Per il mobile TV chiedo la procedura di collaudo antiribaltamento MAIL-31.\r\n')
    doc = upload(env, 'lago', 'richiesta.eml', eml).json()['document']
    p = doc['proposte']
    assert p['tipo'] == 'mail' and p['titolo'] == 'Richiesta collaudo antiribaltamento'
    assert p['data_documento'] == '2026-10-12'
    scheda(env, 'lago', doc['id'], tipo='mail', data_documento=p['data_documento'])
    prompt = ask(env, 'lago')
    assert 'MAIL-31' in prompt and 'mario@esempio.it' in prompt
    monkeypatch.setattr(env.docs, 'oggi', lambda: dt.date(2028, 10, 13))   # oltre 24 mesi
    assert 'MAIL-31' not in ask(env, 'lago')


# ---------------------------------------------------------------- riservato e WhatsApp

def test_riservato_mai_su_whatsapp(env):
    carica(env, 'lago', 'interno.txt', FRASE_LAGO.encode(), tipo='manuale', riservato=True)
    assert 'ZX-77' in ask(env, 'lago')                                  # web: si'
    env.llm.calls.clear()
    asyncio.run(env.core._whatsapp_document_answer('lago', DOMANDA))
    text = '\n'.join(c['text'] for c in env.llm.calls)
    assert text and 'ZX-77' not in text                                 # WhatsApp: mai
    assert env.docs.canale() == 'web'                                   # canale ripristinato


def test_whatsapp_usa_i_documenti_della_sua_azienda(env):
    carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode(), tipo='manuale')
    carica(env, 'covolo', 'covolo.txt', FRASE_COVOLO.encode(), tipo='manuale')
    env.llm.calls.clear()
    asyncio.run(env.core._whatsapp_document_answer('lago', DOMANDA))
    text = '\n'.join(c['text'] for c in env.llm.calls)
    assert 'ZX-77' in text and 'CVZ-SEGRETA' not in text


# ---------------------------------------------------------------- separazione aziende

def test_aziende_separate(env):
    carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode())
    carica(env, 'covolo', 'covolo.txt', FRASE_COVOLO.encode())
    lago, covolo = ask(env, 'lago'), ask(env, 'covolo')
    assert 'ZX-77' in lago and 'CVZ-SEGRETA' not in lago
    assert 'CVZ-SEGRETA' in covolo and 'ZX-77' not in covolo
    elenco = env.http.get('/api/documents/lago', headers=H('lago')).json()
    assert [d['name'] for d in elenco['documents']] == ['lago.txt']
    assert elenco['knowledge_client'] == 'lago'


def test_tenant_senza_archivio_rifiutato_poi_mappa_esplicita(env, monkeypatch):
    assert upload(env, 'roberto', 'x.txt', FRASE_LAGO.encode()).status_code == 403
    monkeypatch.setenv('APPLAUDI_TENANT_CLIENT_MAP', json.dumps({'roberto': 'lago'}))
    r = upload(env, 'roberto', 'x.txt', FRASE_LAGO.encode())
    assert r.status_code == 200 and r.json()['knowledge_client'] == 'lago'
    scheda(env, 'roberto', r.json()['document']['id'], tipo='manuale')
    assert 'ZX-77' in ask(env, 'lago')


def test_chiave_sbagliata_o_di_un_altra_azienda(env):
    assert upload(env, 'lago', 'a.txt', b'x', key='sbagliata').status_code == 401
    assert upload(env, 'lago', 'a.txt', b'x', key=KEYS['covolo']).status_code == 401
    assert env.http.get('/api/documents/lago').status_code == 401
    doc = carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode())
    r = env.http.post(f"/api/documents/covolo/{doc['id']}/scheda", json={'titolo': 'x'}, headers=H('covolo'))
    assert r.status_code == 404                     # covolo non vede la scheda di lago
    r = env.http.get('/api/documents/lago/registro/risposte', headers=H('covolo'))
    assert r.status_code == 401


def test_il_browser_non_sceglie_l_azienda(env):
    r = env.http.post('/api/documents/lago/upload?cliente=covolo',
                      files={'file': ('lago.txt', io.BytesIO(FRASE_LAGO.encode()), 'text/plain')},
                      data={'cliente': 'covolo', 'tenant': 'covolo'}, headers=H('lago'))
    assert r.status_code == 200 and r.json()['knowledge_client'] == 'lago'
    scheda(env, 'lago', r.json()['document']['id'], tipo='manuale')
    assert 'ZX-77' not in ask(env, 'covolo')


# ---------------------------------------------------------------- formati e limiti

def test_formati_rifiutati_e_limiti(env, monkeypatch):
    assert upload(env, 'lago', 'virus.exe', b'MZ').status_code == 415
    assert upload(env, 'lago', 'vuoto.txt', b'   ').status_code == 422
    monkeypatch.setenv('DOCUMENT_MAX_MB', '1')
    assert upload(env, 'lago', 'grande.txt', b'a' * (1024 * 1024 + 10)).status_code == 413
    assert upload(env, 'lago', 'rotto.pdf', b'non sono un pdf').status_code == 422
    files = [p for p in env.root.rglob('*') if p.is_file()]
    assert not [p for p in files if p.name != 'manifest.json']   # nessun residuo


def test_nome_file_ripulito(env):
    r = upload(env, 'lago', '../../etc/passwd.txt', FRASE_LAGO.encode())
    assert r.status_code == 200 and r.json()['document']['name'] == 'passwd.txt'
    assert all(str(p).startswith(str(env.root)) for p in env.root.rglob('*'))


def _pdf_con_testo(testo):
    stream = f'BT /F1 12 Tf 72 720 Td ({testo}) Tj ET'.encode()
    objs = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
            b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R '
            b'/Resources << /Font << /F1 5 0 R >> >> >>',
            b'<< /Length %d >>\nstream\n' % len(stream) + stream + b'\nendstream',
            b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    out, offs = io.BytesIO(), []
    out.write(b'%PDF-1.4\n')
    for i, o in enumerate(objs, 1):
        offs.append(out.tell())
        out.write(b'%d 0 obj\n' % i + o + b'\nendobj\n')
    xref = out.tell()
    out.write(b'xref\n0 %d\n0000000000 65535 f \n' % (len(objs) + 1))
    for o in offs:
        out.write(b'%010d 00000 n \n' % o)
    out.write(b'trailer << /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF' % (len(objs) + 1, xref))
    return out.getvalue()


def test_formati_letti(env):
    from docx import Document
    from openpyxl import Workbook
    doc = Document()
    doc.add_paragraph('Manuale qualita: tolleranza antiribaltamento 3 mm.')
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = 'Collaudo', 'DOCX-TAB-1'
    b = io.BytesIO(); doc.save(b)
    wb = Workbook(); wb.active.append(['Codice', 'Collaudo']); wb.active.append(['XLS-99', 'antiribaltamento'])
    x = io.BytesIO(); wb.save(x)
    casi = {'manuale.docx': (b.getvalue(), 'DOCX-TAB-1'),
            'tabella.xlsx': (x.getvalue(), 'XLS-99'),
            'listino.csv': (b'codice;collaudo\nCSV-42;antiribaltamento\n', 'CSV-42'),
            'note.md': (b'# Collaudo\nantiribaltamento MD-7', 'MD-7'),
            'dati.json': (json.dumps({'collaudo': 'antiribaltamento JSON-8'}).encode(), 'JSON-8'),
            'scheda.pdf': (_pdf_con_testo('Collaudo antiribaltamento PDF-5'), 'PDF-5')}
    for name, (data, _) in casi.items():
        carica(env, 'lago', name, data, tipo='manuale')
    prompt = ask(env, 'lago')
    for name, (_, segno) in casi.items():
        assert segno in prompt, name


# ---------------------------------------------------------------- eliminazione e riavvio

def togli(env, tenant, doc_id, motivo='errore_caricamento', nota='', key=None):
    return env.http.post(f'/api/documents/{tenant}/{doc_id}/togli', json={'motivo': motivo, 'nota': nota},
                         headers={'X-Applaudi-Tenant-Key': key or KEYS[tenant]})


def test_togliere_esce_subito_dalle_risposte_e_si_ripristina(env):
    doc = carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode())
    assert 'ZX-77' in ask(env, 'lago')
    assert togli(env, 'covolo', doc['id']).status_code == 404                # altra azienda
    assert togli(env, 'lago', doc['id'], motivo='').status_code == 400        # motivo obbligatorio
    assert togli(env, 'lago', doc['id'], motivo='altro').status_code == 400   # «altro» vuole la nota
    assert togli(env, 'lago', doc['id']).status_code == 200
    assert 'ZX-77' not in ask(env, 'lago')                                    # fuori SUBITO
    cartella = env.root / 'lago' / 'files' / doc['id']
    assert cartella.exists()                                                  # ma il file resta
    lista = env.http.get('/api/documents/lago?cestino=true', headers=H('lago')).json()
    assert [d['id'] for d in lista['documents']] == [doc['id']] and lista['nel_cestino'] == 1
    assert doc['id'] not in [d['id'] for d in env.http.get('/api/documents/lago', headers=H('lago')).json()['documents']]
    assert env.http.post(f"/api/documents/lago/{doc['id']}/ripristina", headers=H('lago')).status_code == 200
    assert 'ZX-77' in ask(env, 'lago')
    assert togli(env, 'lago', '..%2F..%2Fx').status_code in (400, 404)
    # la vecchia rotta di cancellazione immediata non esiste piu'
    assert env.http.delete(f"/api/documents/lago/{doc['id']}", headers=H('lago')).status_code in (404, 405)
    azioni = [r['azione'] for r in env.http.get('/api/documents/lago/registro/operazioni',
                                                 headers=H('lago')).json()['righe']]
    assert azioni == ['ripristina', 'togli']


def test_freno_contro_le_rimozioni_a_raffica(env, monkeypatch):
    monkeypatch.setenv('RIMOZIONI_ORA', '3')
    docs = [carica(env, 'lago', f'd{n}.txt', f'Documento numero {n} ZX-{n}'.encode()) for n in range(5)]
    assert [togli(env, 'lago', d['id']).status_code for d in docs] == [200, 200, 200, 429, 429]
    reg = env.http.get('/api/documents/lago/registro/operazioni', headers=H('lago')).json()
    assert reg['blocchi_24h'] == 2 and reg['rimozioni_ultima_ora'] == 3
    restano = env.http.get('/api/documents/lago', headers=H('lago')).json()['documents']
    assert {d['id'] for d in restano} == {docs[3]['id'], docs[4]['id']}


def test_distruggere_serve_pin_e_tempo_di_cestino(env, monkeypatch):
    pins = []
    monkeypatch.setattr(env.docs, 'verifica_pin', lambda t, p: pins.append((t, p)) or p == '482913')
    doc = carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode())
    url = f"/api/documents/lago/{doc['id']}/distruggi"
    assert env.http.post(url, json={'pin': '482913'}, headers=H('lago')).status_code == 409   # non nel cestino
    togli(env, 'lago', doc['id'], motivo='superato')
    r = env.http.post(url, json={'pin': '482913'}, headers=H('lago'))
    assert r.status_code == 409 and '30 giorni' in r.json()['detail']                     # troppo presto
    monkeypatch.setenv('CESTINO_GIORNI', '0')
    assert env.http.post(url, json={'pin': '000000'}, headers=H('lago')).status_code == 403   # PIN sbagliato
    assert (env.root / 'lago' / 'files' / doc['id']).exists()
    assert env.http.post(url, json={'pin': '482913'}, headers=H('lago')).status_code == 200
    assert not (env.root / 'lago' / 'files' / doc['id']).exists()
    assert pins[-1] == ('lago', '482913')
    righe = env.http.get('/api/documents/lago/registro/operazioni', headers=H('lago')).json()['righe']
    assert righe[0]['azione'] == 'distruggi' and righe[0]['titolo']


def test_richiesta_privacy_si_distrugge_subito_senza_titolo(env, monkeypatch):
    monkeypatch.setattr(env.docs, 'verifica_pin', lambda t, p: True)
    doc = carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode(), titolo='Dati del signor Rossi')
    togli(env, 'lago', doc['id'], motivo='privacy')
    assert env.http.post(f"/api/documents/lago/{doc['id']}/distruggi", json={'pin': '1'},
                         headers=H('lago')).status_code == 200
    righe = env.http.get('/api/documents/lago/registro/operazioni', headers=H('lago')).json()['righe']
    assert righe[0]['azione'] == 'distruggi' and righe[0]['titolo'] is None
    assert 'Rossi' not in json.dumps(righe[0])


def test_dopo_un_riavvio_i_documenti_restano(env):
    core = env.core
    carica(env, 'lago', 'lago.txt', FRASE_LAGO.encode())
    core.attiva_cliente('covolo')
    core.STATO_PER_CLIENTE.pop('lago', None)     # come un servizio appena ripartito
    core.attiva_cliente('lago')
    assert any('ZX-77' in p['text'] for p in core.UPLOADED_PAGES)
    assert 'ZX-77' in ask(env, 'lago')


# ---------------------------------------------------------------- chi dice cosa

RISPOSTA_CON_SEZIONE = (
    "RISPOSTA\n"
    "Per un mobile TV largo circa 160 cm il riferimento e' N.O.W. TV UNITS codice 211605: "
    "larghezza 163,8 cm, prezzo 2.840 nella finitura vetro (pag. 60).\n\n"
    "DAI DOCUMENTI AZIENDALI\n"
    "- [DIVERSO DAL CATALOGO] prezzo promozionale 2.500 per il codice 211605, contro 2.840 "
    "di listino (pag. 60) — fonte: «Promo Autunno» (valido dal 01/10/2026 al 30/11/2026), documento\n"
    "- [NUOVO] solo finitura vetro opaco — fonte: «Promo Autunno», documento")


def ask_answer(env, question=DOMANDA_PROMO):
    env.core.ANSWER_CACHE.clear()
    r = env.http.post('/api/ask', json={'question': '/catalogo ' + question, 'cliente': 'lago'})
    assert r.status_code == 200, r.text
    return r.json()['answer']


def test_sezione_documenti_arriva_intatta(env, monkeypatch):
    llm = FakeLLM(env.core, answer=RISPOSTA_CON_SEZIONE)
    monkeypatch.setattr(env.core, 'client', llm)
    carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione', titolo='Promo Autunno',
           valido_dal='2026-10-01', valido_fino='2026-11-30')
    answer = ask_answer(env)
    assert 'DAI DOCUMENTI AZIENDALI' in answer
    principale, sezione = answer.split('DAI DOCUMENTI AZIENDALI', 1)
    assert '2.500' in sezione and '«Promo Autunno»' in sezione and '30/11/2026' in sezione
    assert '2.840' in principale and '2.500' not in principale   # non fusi
    prompt = [c['text'] for c in llm.calls if 'RICHIESTA ORIGINALE' in c['text']][0]
    assert '[DIVERSO DAL CATALOGO]' in prompt and "oggi e' il 17/10/2026" in prompt


def test_sezione_inventata_senza_documenti_viene_tolta(env, monkeypatch):
    monkeypatch.setattr(env.core, 'client', FakeLLM(env.core, answer=RISPOSTA_CON_SEZIONE))
    answer = ask_answer(env)                       # nessun documento caricato
    assert 'DAI DOCUMENTI AZIENDALI' not in answer and '2.500' not in answer


def test_whatsapp_simboli_al_posto_dei_colori(env, monkeypatch):
    monkeypatch.setattr(env.core, 'client', FakeLLM(env.core, answer=RISPOSTA_CON_SEZIONE))
    carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione', titolo='Promo Autunno',
           valido_dal='2026-10-01', valido_fino='2026-11-30')
    env.core.ANSWER_CACHE.clear()
    testo = asyncio.run(env.core._whatsapp_document_answer('lago', DOMANDA_PROMO))
    assert '*DAI DOCUMENTI AZIENDALI*' in testo
    assert '🟠 *DIVERSO DAL CATALOGO*' in testo and '🟢 *NUOVO*' in testo


def test_split_section():
    import documenti_aziendali as d
    head, sez = d.split_section('A\n\n**Dai documenti aziendali:**\n- [NUOVO] x')
    assert head == 'A' and sez == 'DAI DOCUMENTI AZIENDALI\n- [NUOVO] x'
    assert d.split_section('nessuna sezione') == ('nessuna sezione', '')
    assert d.split_section('A\nDAI DOCUMENTI AZIENDALI\n') == ('A', '')


# ---------------------------------------------------------------- registro per le contestazioni

def test_registro_risposte_con_fonti(env, monkeypatch):
    monkeypatch.setattr(env.core, 'client', FakeLLM(env.core, answer=RISPOSTA_CON_SEZIONE))
    doc = carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione', titolo='Promo Autunno',
                 valido_dal='2026-10-01', valido_fino='2026-11-30')
    ask_answer(env)
    asyncio.run(env.core._whatsapp_document_answer('lago', DOMANDA_PROMO))
    righe = env.http.get('/api/documents/lago/registro/risposte?testo=211605',
                         headers=H('lago')).json()['righe']
    assert len(righe) == 2
    wa, web = righe                                   # dal piu' recente
    assert wa['canale'] == 'whatsapp' and web['canale'] == 'web'
    assert 60 in web['pagine_catalogo'] and web['sezione_documenti']
    assert web['documenti'] == [{'id': doc['id'], 'titolo': 'Promo Autunno', 'versione': doc['versione']}]
    assert wa['documenti'] and wa['documenti'][0]['id'] == doc['id']


def test_dashboard_pagina_documenti(env):
    html = env.http.get('/plaud/lago').text
    sezione = html.split('id="documents"')[1].split('id="proposals"')[0]
    assert 'Carica e leggi' in sezione and 'Non operativo' not in sezione
    assert 'Conferma e inserisci' in sezione and 'Registro delle risposte' in sezione


def test_nessuna_contraddizione_con_la_sezione(env, monkeypatch):
    risposta = ("RISPOSTA\nCodice 211605 a 2.840 (pag. 60).\nPromozioni: nel catalogo non e' "
                "indicata alcuna promozione. Informazione non trovata nel documento collegato.\n\n"
                "DAI DOCUMENTI AZIENDALI\n- [DIVERSO DAL CATALOGO] 211605 a 2.490 contro 2.840 "
                "(pag. 60) — fonte: «Promo Autunno»\n- [NUOVO] consegna in 15 giorni — fonte: «Promo Autunno»")
    llm = FakeLLM(env.core, answer=risposta)
    monkeypatch.setattr(env.core, 'client', llm)
    carica(env, 'lago', 'promo.txt', PROMO.encode(), tipo='promozione', titolo='Promo Autunno',
           valido_dal='2026-10-01', valido_fino='2026-11-30')
    answer = ask_answer(env)
    principale = answer.split('\n\nDAI DOCUMENTI AZIENDALI', 1)[0]
    assert 'Informazione non trovata' not in principale and '2.840' in principale
    prompt = [c['text'] for c in llm.calls if 'RICHIESTA ORIGINALE' in c['text']][0]
    assert "non scrivere MAI 'informazione non trovata'" in prompt and 'SEMPRE [DIVERSO DAL' in prompt
    assert 'DIRETTO e BREVE' in prompt


def test_senza_sezione_il_non_trovata_resta(env, monkeypatch):
    monkeypatch.setattr(env.core, 'client', FakeLLM(
        env.core, answer='RISPOSTA\nInformazione non trovata nel documento collegato.'))
    assert 'Informazione non trovata' in ask_answer(env)


def test_varianti_del_non_trovata_tolte():
    import documenti_aziendali as d
    head = ("RISPOSTA\nNel catalogo non c'e'. Informazione non trovata nel documento collegato per "
            "quanto riguarda promozioni a catalogo.\nAltro testo.")
    out = d.join_section(head, 'DAI DOCUMENTI AZIENDALI\n- [NUOVO] x')
    assert 'non trovata' not in out and 'Altro testo.' in out and out.endswith('- [NUOVO] x')


def test_codici_del_documento_portano_la_riga_di_catalogo(env):
    """Domanda senza codici ("abbiamo promozioni?"): il codice nominato dal documento,
    anche in una tabella senza la parola 'codice', deve far leggere la sua riga di listino."""
    from docx import Document
    doc = Document()
    doc.add_paragraph('Promozione valida fino al 30/11/2026.')
    t = doc.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ['Articolo', 'Finitura', 'Prezzo promo']): c.text = v
    for c, v in zip(t.rows[1].cells, ['211605', 'vetro opaco', '2.490']): c.text = v
    b = io.BytesIO(); doc.save(b)
    carica(env, 'lago', 'promo.docx', b.getvalue(), tipo='promozione', titolo='Promo Autunno',
           valido_dal='2026-10-01', valido_fino='2026-11-30')
    prompt = ask(env, 'lago', 'Abbiamo qualche promozione al momento?')
    assert '2.490' in prompt                               # il documento c'e'
    assert '- 211605 | pagina 60' in prompt and '2.840' in prompt   # e la riga di listino anche


def test_senza_documenti_nessuna_riga_in_piu(env):
    prompt = ask(env, 'lago', 'Abbiamo qualche promozione al momento?')
    assert '- 211605 | pagina 60' not in prompt


# ---------------------------------------------------------------- significato e caso Manzoni

ROMANZO = ("I PROMESSI SPOSI di Alessandro Manzoni. Quel ramo del lago di Como, che volge a "
           "mezzogiorno. Renzo e Lucia, promessi sposi, sono ostacolati da don Rodrigo. "
           "Il Manzoni racconta la peste di Milano. Manzoni, Manzoni, Manzoni: il romanzo "
           "del Manzoni e' il capolavoro di Manzoni.")
LISTINO_MANZONI = ("LISTINO RUBINETTERIA 2026. Linea Manzoni: rubinetto Manzoni miscelatore lavabo, "
                   "codice MZ-1001, finitura cromo, prezzo 320 euro.")


class FakeEmbedder:
    """Significato finto ma coerente: vettore delle parole (hash in 64 dimensioni)."""

    def __init__(self):
        self.calls = []

    def __call__(self, testi):
        import re as _re
        self.calls.append(list(testi))
        out = []
        for t in testi:
            v = [0.0] * 64
            for w in _re.findall(r'[a-zàèéìòù]{3,}', t.lower()):
                v[hash(w) % 64] += 1.0
            out.append(v)
        return out


@pytest.fixture
def emb(env, monkeypatch):
    f = FakeEmbedder()
    monkeypatch.setattr(env.docs, 'embedder', f)
    env.docs._QUERY_CACHE.clear()
    return f


def test_manzoni_domanda_di_prezzo_mai_dal_romanzo(env, emb):
    carica(env, 'lago', 'promessi_sposi.txt', ROMANZO.encode(), tipo='altro', titolo='I promessi sposi')
    carica(env, 'lago', 'listino_manzoni.txt', LISTINO_MANZONI.encode(), tipo='listino',
           titolo='Listino rubinetteria 2026')
    prompt = ask(env, 'lago', 'Quanto costa ora il rubinetto Manzoni?')
    assert 'MZ-1001' in prompt and '320 euro' in prompt
    assert 'Renzo' not in prompt and 'promessi sposi' not in prompt.lower()


def test_manzoni_domanda_sul_romanzo_trova_il_romanzo(env, emb):
    carica(env, 'lago', 'promessi_sposi.txt', ROMANZO.encode(), tipo='altro', titolo='I promessi sposi')
    carica(env, 'lago', 'listino_manzoni.txt', LISTINO_MANZONI.encode(), tipo='listino',
           titolo='Listino rubinetteria 2026')
    prompt = ask(env, 'lago', 'Chi sono Renzo e Lucia nel romanzo del Manzoni?')
    assert 'Renzo e Lucia' in prompt


def test_una_sola_parola_in_comune_non_basta(env):
    carica(env, 'lago', 'ferie.txt', b'Calendario ferie estive del magazzino di Bassano.', tipo='altro')
    prompt = ask(env, 'lago', 'Che tempo fa a Bassano oggi?')
    assert 'Calendario ferie' not in prompt


def test_codice_alla_lettera_basta(env):
    carica(env, 'lago', 'listino.txt', LISTINO_MANZONI.encode(), tipo='listino')
    assert 'MZ-1001' in ask(env, 'lago', 'MZ-1001?')


def test_significato_calcolato_alla_conferma_salvato_e_riletto(env, emb):
    doc = carica(env, 'lago', 'listino.txt', LISTINO_MANZONI.encode(), tipo='listino')
    assert doc['significato'] is True and len(emb.calls) == 1
    env.core.attiva_cliente('covolo')
    env.core.STATO_PER_CLIENTE.pop('lago', None)          # riavvio: i vettori si rileggono
    env.core.attiva_cliente('lago')
    assert any(p['vec'] for p in env.core.UPLOADED_PAGES)
    assert len(emb.calls) == 1                             # nessun ricalcolo al riavvio


def test_riservato_mai_al_servizio_del_significato(env, emb):
    doc = carica(env, 'lago', 'interno.txt', b'Margini interni riservati MARGINE-77 rubinetto.',
                 tipo='listino', riservato=True)
    assert doc['significato'] is False
    assert not any('MARGINE-77' in t for call in emb.calls for t in call)
    d = scheda(env, 'lago', doc['id'], riservato=False).json()['document']      # tolto il riservato
    assert d['significato'] is True
    d = scheda(env, 'lago', doc['id'], riservato=True).json()['document']       # rimesso
    assert d['significato'] is False


def test_servizio_del_significato_giu_il_documento_funziona(env, monkeypatch):
    def rotto(testi):
        raise RuntimeError('servizio non raggiungibile')
    monkeypatch.setattr(env.docs, 'embedder', rotto)
    env.docs._QUERY_CACHE.clear()
    doc = carica(env, 'lago', 'collaudo.txt', FRASE_LAGO.encode(), tipo='manuale')
    assert doc['stato'] == 'attivo' and doc['significato'] is False
    assert 'ZX-77' in ask(env, 'lago')                    # si cerca per parole


# ---------------------------------------------------------------- prezzo <-> colonna
# Collaudo 19ott2026: sul web la risposta diceva "listino in vetro opaco 3.958" (e' la
# colonna XGlass del 211605, pag. 60) e calcolava uno sconto del 37% invece del 12%.
SBAGLIATA = ("RISPOSTA\nLa promozione riguarda il codice 211605 N.O.W. TV UNITS in vetro opaco a "
             "2.490 euro. Il prezzo di listino dello stesso codice in vetro opaco e' 3.958 "
             "(pag. 60), quindi lo sconto e' di circa il 37%.")
GIUSTA = ("RISPOSTA\nIl codice 211605 N.O.W. TV UNITS in vetro opaco e' in promozione a 2.490 euro; "
          "il listino in vetro opaco e' 2.840 (pag. 60). La finitura XGlass resta a listino: 3.958.")


class StesuraInSequenza(FakeLLM):
    """Le stesure (messaggi con RICHIESTA ORIGINALE) rispondono in ordine dalla lista."""

    def __init__(self, core, stesure):
        super().__init__(core, answer='')
        self.stesure = list(stesure)
        originale = self.chat.completions.create

        def create(**kw):
            text = '\n'.join(m['content'] for m in kw['messages'])
            if 'RICHIESTA ORIGINALE' not in text:
                return originale(**kw)
            self.calls.append({'cliente': core.CLIENTE, 'text': text})
            content = self.stesure.pop(0) if len(self.stesure) > 1 else self.stesure[0]
            msg = types.SimpleNamespace(content=content)
            return types.SimpleNamespace(choices=[types.SimpleNamespace(message=msg)])
        self.chat.completions.create = create


def stesure(llm):
    return [c['text'] for c in llm.calls if 'RICHIESTA ORIGINALE' in c['text']]


def test_controllo_riconosce_prezzo_nella_colonna_sbagliata(env):
    errori = env.core.verifica_prezzi_colonne(SBAGLIATA)
    assert [(e['codice'], e['prezzo']) for e in errori] == [('211605', '3.958')]
    assert env.core.verifica_prezzi_colonne(GIUSTA) == []
    # nessun allarme se la frase non nomina una finitura o le nomina entrambe
    assert env.core.verifica_prezzi_colonne('Il 211605 costa 2.840 euro.') == []
    assert env.core.verifica_prezzi_colonne(
        'Codice 211605: vetro opaco o XGlass, rispettivamente 2.840 e 3.958 euro.') == []


def test_prezzo_sbagliato_si_riscrive_con_gli_abbinamenti_esatti(env, monkeypatch):
    llm = StesuraInSequenza(env.core, [SBAGLIATA, GIUSTA])
    monkeypatch.setattr(env.core, 'client', llm)
    answer = ask_answer(env)
    assert '2.840' in answer and '37%' not in answer
    seconda = stesure(llm)[1]
    assert 'CORREZIONE OBBLIGATORIA' in seconda and '2.840' in seconda and '3.958' in seconda


def test_prezzo_sbagliato_due_volte_escono_solo_i_prezzi_certi(env, monkeypatch):
    llm = StesuraInSequenza(env.core, [SBAGLIATA, SBAGLIATA])
    monkeypatch.setattr(env.core, 'client', llm)
    answer = ask_answer(env)
    assert '37%' not in answer and '2.490' not in answer
    assert '2.840' in answer and '3.958' in answer and 'XGlass' in answer


def test_risposta_giusta_nessuna_stesura_in_piu(env, monkeypatch):
    llm = StesuraInSequenza(env.core, [GIUSTA])
    monkeypatch.setattr(env.core, 'client', llm)
    assert '2.840' in ask_answer(env)
    assert len(stesure(llm)) == 1


def test_riscrittura_ancora_sbagliata_non_passa(env, monkeypatch):
    # il controllo dopo la stesura, da solo: la seconda stesura sbagliata non viene accettata
    monkeypatch.setattr(env.core, 'generation_chat', lambda *a, **k: SBAGLIATA)
    env.core._SEZIONE_DOCUMENTI.testo = ''
    out = env.core._controllo_prezzi_colonne(SBAGLIATA, 'sistema', 'domanda', 500, False)
    assert '37%' not in out and '3.958' in out and '2.840' in out
    monkeypatch.setattr(env.core, 'generation_chat', lambda *a, **k: GIUSTA)
    assert env.core._controllo_prezzi_colonne(SBAGLIATA, 'sistema', 'domanda', 500, False) == GIUSTA


def test_errore_introdotto_dai_controlli_successivi_non_esce(env, monkeypatch):
    # stesura giusta, ma un controllo successivo la riscrive sbagliata: l'ultima verifica la ferma
    monkeypatch.setattr(env.core, 'client', StesuraInSequenza(env.core, [GIUSTA]))
    monkeypatch.setattr(env.core, 'apply_contradiction_control', lambda answer, hits: SBAGLIATA)
    answer = ask_answer(env)
    assert '37%' not in answer and '2.840' in answer and '3.958' in answer


# ---------------------------------------------------------------- Narratore dalla dashboard
# Collaudo 19ott2026: "Analizza" dalla dashboard rispondeva "Integrazione Plaud non
# configurata sul server" perche' passava dal token della rotta esterna, assente sul
# servizio di prova. La dashboard entra con la chiave del cliente e basta.
ANALISI = json.dumps({"tipo_comunicazione": "Riunione", "richiesta_cliente": "Ordine mobile TV",
                      "sintesi_operativa": "Ordine confermato", "instradamento": "CONVERSAZIONE",
                      "rilevanza_catalogo": False, "domanda_catalogo": ""})


def test_analizza_dalla_dashboard_senza_token_esterno(env, monkeypatch):
    monkeypatch.setattr(env.core, 'PLAUD_INGEST_TOKEN', '')
    monkeypatch.setattr(env.core, 'client', FakeLLM(env.core, answer=ANALISI))

    async def bundle(slot, rid):
        assert slot == 'roberto#privato' and rid == 'abc123'
        return {"transcript": "Ordiniamo il mobile TV Now.", "note": "", "detail": {"name": "Riunione"}}
    monkeypatch.setattr(env.core, 'get_recording_bundle', bundle)
    monkeypatch.setattr(env.core, 'plaud_slot', lambda t, f: t + ('#privato' if f == 'privato' else ''))
    r = env.http.post('/api/plaud/cloud/analyze/roberto/abc123?fonte=privato', headers=H('roberto'))
    assert r.status_code == 200, r.text
    assert r.json()['analysis']['richiesta_cliente'] == 'Ordine mobile TV'
    # senza chiave del cliente non si entra
    assert env.http.post('/api/plaud/cloud/analyze/roberto/abc123?fonte=privato').status_code in (401, 403)


def test_rotta_esterna_resta_protetta_dal_token(env, monkeypatch):
    corpo = {"recording_id": "of_" + "a" * 32, "transcript": "testo"}
    monkeypatch.setattr(env.core, 'PLAUD_INGEST_TOKEN', '')
    assert env.http.post('/api/plaud/analyze', json=corpo).status_code == 503
    monkeypatch.setattr(env.core, 'PLAUD_INGEST_TOKEN', 'giusto')
    assert env.http.post('/api/plaud/analyze', json=corpo,
                         headers={'X-Plaud-Token': 'sbagliato'}).status_code == 401
    assert env.http.post('/api/plaud/analyze', json=dict(corpo, recording_id='abc'),
                         headers={'X-Plaud-Token': 'giusto'}).status_code == 400


class AnalisiInSequenza(FakeLLM):
    def __init__(self, core, risposte):
        super().__init__(core, answer='')
        self.risposte, self.budget = list(risposte), []

        def create(**kw):
            self.budget.append(kw.get('max_tokens') or kw.get('max_completion_tokens'))
            msg = types.SimpleNamespace(content=self.risposte.pop(0))
            return types.SimpleNamespace(choices=[types.SimpleNamespace(message=msg)])
        self.chat.completions.create = create


def analizza(env, monkeypatch, risposte):
    llm = AnalisiInSequenza(env.core, risposte)
    monkeypatch.setattr(env.core, 'client', llm)

    async def bundle(slot, rid):
        return {"transcript": "Ordiniamo il mobile TV Now.", "note": "", "detail": {"name": "Riunione"}}
    monkeypatch.setattr(env.core, 'get_recording_bundle', bundle)
    monkeypatch.setattr(env.core, 'plaud_slot', lambda t, f: t)
    return env.http.post('/api/plaud/cloud/analyze/roberto/abc123', headers=H('roberto')), llm


def test_analisi_tagliata_si_riprova(env, monkeypatch):
    r, llm = analizza(env, monkeypatch, [ANALISI[:60], ANALISI])     # prima: JSON tagliato
    assert r.status_code == 200, r.text
    assert len(llm.budget) == 2 and llm.budget[0] >= 4000


def test_analisi_con_testo_intorno_si_legge(env, monkeypatch):
    r, llm = analizza(env, monkeypatch, ['Ecco l\'analisi:\n```json\n' + ANALISI + '\n```\nFine.'])
    assert r.status_code == 200, r.text
    assert len(llm.budget) == 1


def test_analisi_illeggibile_due_volte_errore_chiaro(env, monkeypatch):
    r, _ = analizza(env, monkeypatch, ['non json', '{"tagliato": '])
    assert r.status_code == 502


# ---------------------------------------------------------------- codici dettati a gruppi
# Collaudo 19ott2026: PLAUD trascrive "codice 211 605" e il motore rispondeva che il
# codice non compare, pur avendo la riga del 211605.
def test_codice_dettato_a_gruppi_si_riconosce(env):
    prompt = ask(env, 'lago', 'Quanto costa il mobile TV Now codice 211 605 in vetro opaco?')
    assert '211605' in prompt and '211 605' not in prompt


def test_unione_solo_per_codici_veri(env):
    unisci = env.core.unisci_codici_spezzati
    assert unisci('codice 211 605 e 21 16 06') == 'codice 211605 e 211606'
    assert unisci('costa 1 400 euro, tavolo 160 x 90') == 'costa 1 400 euro, tavolo 160 x 90'
    assert unisci('codici 211605 211606') == 'codici 211605 211606'
    assert unisci('ne servono 2 11605') == 'ne servono 2 11605'      # pezzo di 1 cifra: un numero


def test_video_nel_riquadro_documenti_spiega_dove_va(env):
    r = upload(env, 'lago', 'sopralluogo.mp4', b'\x00\x00\x00\x18ftypmp42' + b'0' * 100)
    assert r.status_code == 415 and 'trigger' in r.json()['detail']


def test_tabella_lunga_ogni_parte_ha_l_intestazione(tmp_path):
    import documenti_aziendali as docs
    righe = ['Codice;Cliente;Indirizzo;Telefono'] + [
        f'CL-{n:04d};Cliente numero {n};Via delle Prove {n}, Bassano;0424 000 {n:03d}' for n in range(400)]
    f = tmp_path / 'anagrafica.csv'
    f.write_text('\n'.join(righe), encoding='utf-8')
    parti, _ = docs.extract(str(f), '.csv')
    assert len(parti) > 1
    for p in parti:
        testo = p['text'].split('\n')
        assert testo[0] == 'Codice\tCliente\tIndirizzo\tTelefono'          # intestazione in ogni parte
        assert all(r.startswith('CL-') and r.count('\t') == 3 for r in testo[1:])   # righe intere
        assert len(p['text']) <= docs.EVIDENCE_CHARS + 100
    tutte = [r for p in parti for r in p['text'].split('\n')[1:]]
    assert len(tutte) == 400                                                # nessuna riga persa


# ------------------------------------------------------------------ domande senza risposta -> FAQ
def test_esito_della_risposta():
    import documenti_aziendali as docs
    assert docs.esito_risposta("Informazione non trovata nel documento collegato.") == "senza_risposta"
    assert docs.esito_risposta("Non ho trovato prove sufficienti per rispondere.") == "senza_risposta"
    assert docs.esito_risposta("Il codice costa 2.840 euro. La garanzia: informazione non trovata.") == "debole"
    assert docs.esito_risposta("Il codice 211605 costa 2.840 euro (pag. 60).") == "risposta"
    assert docs.esito_risposta("qualunque cosa", fallita=True) == "senza_risposta"


def test_domande_senza_risposta_diventano_faq(env, monkeypatch):
    monkeypatch.setattr(env.core, 'client', FakeLLM(env.core, answer='Informazione non trovata nel documento collegato.'))
    ask_answer(env, 'Fate la permuta della bici vecchia?')
    ask_answer(env, 'fate la permuta della bici vecchia')
    monkeypatch.setattr(env.core, 'client', FakeLLM(env.core, answer='Si, a noleggio dal lunedi. La tariffa: informazione non trovata.'))
    ask_answer(env, 'Si puo noleggiare una bici?')
    r = env.http.get('/api/documents/lago/senza-risposta', headers=H('lago')).json()
    assert [(g['esito'], g['volte']) for g in r['righe']] == [('senza_risposta', 2), ('debole', 1)]
    assert r['righe'][0]['canali'] == ['web']
    # un altro cliente non vede le domande di lago
    assert env.http.get('/api/documents/covolo/senza-risposta', headers=H('covolo')).json()['righe'] == []
    # con un tocco diventa una FAQ: entra nella conoscenza e la domanda esce dalla lista
    x = env.http.post('/api/documents/lago/faq', json={'domanda': r['righe'][0]['domanda'],
                                                       'risposta': 'Sì, valutiamo la bici usata in negozio.'},
                      headers=H('lago'))
    assert x.status_code == 200, x.text
    doc = x.json()['document']
    assert doc['titolo'].startswith('FAQ: ') and doc['stato'] == 'attivo'
    elenco = env.http.get('/api/documents/lago', headers=H('lago')).json()['documents']
    assert any(d['id'] == doc['id'] for d in elenco)
    r = env.http.get('/api/documents/lago/senza-risposta', headers=H('lago')).json()
    assert [g['esito'] for g in r['righe']] == ['debole']
    # la prossima volta la FAQ arriva al motore come fonte
    llm = FakeLLM(env.core, answer='Sì, valutiamo la bici usata in negozio.')
    monkeypatch.setattr(env.core, 'client', llm)
    ask_answer(env, 'Fate la permuta della bici vecchia?')
    assert 'valutiamo la bici usata' in json.dumps(llm.calls, ensure_ascii=False)
    # ignorata: non torna
    env.http.post('/api/documents/lago/senza-risposta/ignora', json={'domanda': 'Si puo noleggiare una bici?'},
                  headers=H('lago'))
    assert env.http.get('/api/documents/lago/senza-risposta', headers=H('lago')).json()['righe'] == []
    # la stessa FAQ non entra due volte
    assert env.http.post('/api/documents/lago/faq', json={'domanda': 'Fate la permuta della bici vecchia?',
                                                          'risposta': 'altro'}, headers=H('lago')).status_code == 409
    assert env.http.post('/api/documents/lago/faq', json={'domanda': 'x', 'risposta': 'y'},
                         headers={'X-Applaudi-Tenant-Key': 'sbagliata'}).status_code in (401, 403)


def test_domanda_fuori_tema_non_diventa_descrizione_dei_listini(env, monkeypatch):
    """23ott2026: «Fate uno sconto a chi ha già la polizza casa?» sull'archivio LAGO ha
    prodotto una descrizione delle colonne dei listini. Ora la richiesta e' ripetuta in fondo,
    con l'ordine di rispondere «Informazione non trovata» se le evidenze non c'entrano; e quella
    risposta finisce tra le domande senza risposta."""
    domanda = "Fate uno sconto a chi ha già la polizza casa con voi?"
    llm = FakeLLM(env.core, answer="Informazione non trovata nei documenti.")
    monkeypatch.setattr(env.core, 'client', llm)
    ask_answer(env, domanda)
    finale = llm.calls[-1]['text']
    assert "RICHIESTA DA SODDISFARE" in finale and "polizza casa" in finale.split("RICHIESTA DA SODDISFARE")[1]
    assert "Informazione non trovata nei documenti." in finale
    elenco = env.http.get('/api/documents/lago/senza-risposta', headers=H('lago')).json()
    assert any("polizza casa" in (r.get('domanda') or '') and r['esito'] == 'senza_risposta'
               for r in elenco['righe'])
